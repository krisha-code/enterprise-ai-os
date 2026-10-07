from docx import Document


file_path = "uploads/documents/fb563b86-579e-4c25-b0c6-5736a8ba492f.docx"

document = Document(file_path)

print(f"Total paragraphs: {len(document.paragraphs)}")

for index, paragraph in enumerate(document.paragraphs[:100]):

    text = paragraph.text.strip()

    if not text:
        continue

    print("\n" + "=" * 70)
    print(f"Paragraph: {index}")
    print(f"Style:     {paragraph.style.name}")
    print(f"Text:      {text[:200]}")